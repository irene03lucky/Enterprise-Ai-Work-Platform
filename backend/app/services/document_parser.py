"""文档解析器：多格式文件 → 纯文本。

支持 PDF / DOCX / PPTX / TXT / Markdown。
文档是企业资产，解析结果是后续分块与向量化的统一输入。
"""

from pathlib import Path

SUPPORTED_EXTENSIONS = {"pdf", "docx", "pptx", "txt", "md"}


def parse_file(path: str | Path, file_type: str) -> str:
    """按类型解析文件并返回纯文本。解析失败抛出 ValueError。"""
    p = Path(path)
    if not p.exists():
        raise ValueError(f"文件不存在: {p}")

    if file_type == "pdf":
        return _parse_pdf(p)
    if file_type == "docx":
        return _parse_docx(p)
    if file_type == "pptx":
        return _parse_pptx(p)
    if file_type in {"txt", "md"}:
        return _parse_text(p)
    raise ValueError(f"不支持的文件类型: {file_type}")


def _parse_pdf(p: Path) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(str(p))
        pages = [page.extract_text() or "" for page in reader.pages]
    except Exception as exc:  # noqa: BLE001
        raise ValueError(f"PDF 解析失败: {exc}") from exc
    return "\n\n".join(text.strip() for text in pages if text.strip())


def _parse_docx(p: Path) -> str:
    import docx

    try:
        document = docx.Document(str(p))
        paragraphs = [para.text.strip() for para in document.paragraphs]
        # 表格中的文本同样纳入
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text.strip())
    except Exception as exc:  # noqa: BLE001
        raise ValueError(f"DOCX 解析失败: {exc}") from exc
    return "\n".join(text for text in paragraphs if text)


def _parse_pptx(p: Path) -> str:
    from pptx import Presentation

    try:
        presentation = Presentation(str(p))
        texts: list[str] = []
        for slide in presentation.slides:
            slide_texts: list[str] = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    content = shape.text_frame.text.strip()
                    if content:
                        slide_texts.append(content)
            if slide_texts:
                texts.append(f"[第 {len(texts) + 1} 页]\n" + "\n".join(slide_texts))
    except Exception as exc:  # noqa: BLE001
        raise ValueError(f"PPTX 解析失败: {exc}") from exc
    return "\n\n".join(texts)


def _parse_text(p: Path) -> str:
    try:
        return p.read_text(encoding="utf-8", errors="ignore")
    except Exception as exc:  # noqa: BLE001
        raise ValueError(f"文本读取失败: {exc}") from exc
